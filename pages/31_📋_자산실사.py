# -*- coding: utf-8 -*-
"""
자산 실사 (휴대폰용) — 찍자마자 AI가 자산번호를 읽어서 자산대장과 대조

사용 흐름
1) 자산대장 파일(.xlsx 또는 .csv)을 올립니다.
   - 회사 보안(DRM)으로 암호화된 엑셀은 서버에서 읽을 수 없으므로,
     보안을 해제하고 외부에 나가도 괜찮은 열만 남긴 파일을 올립니다.
2) [📷 촬영] 버튼을 누르고 자산번호 라벨을 찍으면, Claude AI가 번호를 읽어서
   대장에 있으면 바로 "있음"으로 표시합니다. (확인 버튼 누를 필요 없음)
3) 진행 기록은 서버에 자동 저장되어, 휴대폰 화면이 꺼지거나 새로고침돼도 이어서 할 수 있습니다.
4) [결과 엑셀 저장]으로 "실사결과(있음/미확인)" 열이 붙은 엑셀을 내려받습니다.
"""

import streamlit as st
import pandas as pd
import os
import sys
import re
import json
import base64
import hashlib
from io import BytesIO
from datetime import datetime, timedelta, timezone

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
from style import apply_common_style, show_page_header, check_password
from api_config import get_key, set_key

st.set_page_config(page_title="자산실사 - 마이 유틸리티", page_icon="📋", layout="centered")
apply_common_style()
check_password()
show_page_header("📋", "자산 실사", "자산번호 라벨을 찍으면 AI가 번호를 읽어서 자산대장에 있는지 바로 알려드려요")

# ── 필요한 패키지 확인 ──────────────────────────────────────
try:
    import anthropic
    from PIL import Image, ImageOps
except ImportError:
    st.error("필요한 프로그램(패키지)이 설치되어 있지 않아요. 터미널에서 아래 명령을 실행한 뒤 새로고침 해주세요.")
    st.code("pip install anthropic pillow openpyxl", language="bash")
    st.stop()

# 아이폰 사진(.heic)을 열 수 있게 해주는 부품 (있으면 사용)
try:
    import pillow_heif
    pillow_heif.register_heif_opener()
except ImportError:
    pass

MODEL = "claude-opus-5"
SAVE_DIR = os.path.join(os.path.dirname(__file__), "..", "data", "asset_inspection")   # 진행 기록 저장 폴더
KST = timezone(timedelta(hours=9))   # 서버 시간이 달라도 한국 시간으로 기록

VISION_PROMPT = """이 사진은 회사 자산(비품·장비 등)에 붙어 있는 자산번호 라벨입니다.
사진에서 자산번호(관리번호/일련번호)로 보이는 코드를 찾아주세요.
자산번호는 보통 영문자·숫자·하이픈(-)이 섞인 코드입니다. (예: A-2023-0015, IT00123)

규칙:
- 자산번호로 보이는 코드를 한 줄에 하나씩, 가능성이 높은 것부터 적어주세요. (최대 5개)
- 글자는 사진에 보이는 그대로 적고, 설명·따옴표·번호 매기기는 붙이지 마세요.
- 회사 이름, 전화번호, 날짜처럼 자산번호가 아닌 것은 빼주세요.
- 자산번호로 보이는 것이 전혀 없거나 읽을 수 없으면 NOT_FOUND 한 단어만 적어주세요."""

제목_키워드 = ["자산번호", "일련번호", "자산코드", "관리번호", "번호"]

# ── 세션 상태(이 화면에서만 쓰는 임시 기억) 초기화 ──────────
st.session_state.setdefault("ai_job_id", None)       # 지금 진행 중인 실사 기록 이름
st.session_state.setdefault("ai_photo_key", 0)       # 촬영 버튼을 비우기 위한 번호
st.session_state.setdefault("ai_last", None)         # 마지막 인식 결과


def now_text():
    return datetime.now(KST).strftime("%Y-%m-%d %H:%M:%S")


# ══════════════════════════════════════════════════════════════
# 실사 기록 저장/불러오기 (서버 파일)
# ══════════════════════════════════════════════════════════════
def job_path(job_id):
    return os.path.join(SAVE_DIR, f"{job_id}.json")


def load_job(job_id):
    try:
        with open(job_path(job_id), "r", encoding="utf-8") as f:
            return json.load(f)
    except Exception:
        return None


def save_job(job):
    os.makedirs(SAVE_DIR, exist_ok=True)
    job["updated"] = now_text()
    with open(job_path(job["id"]), "w", encoding="utf-8") as f:
        json.dump(job, f, ensure_ascii=False)


def list_jobs():
    """저장된 실사 기록 목록 (최근 것부터)"""
    if not os.path.isdir(SAVE_DIR):
        return []
    jobs = []
    for name in os.listdir(SAVE_DIR):
        if name.endswith(".json"):
            job = load_job(name[:-5])
            if job:
                jobs.append(job)
    return sorted(jobs, key=lambda j: j.get("updated", ""), reverse=True)


# ══════════════════════════════════════════════════════════════
# 자산대장 파일 읽기
# ══════════════════════════════════════════════════════════════
def rows_to_table(rows):
    """시트 내용을 표로 바꿉니다. 맨 위 제목 줄은 건너뛰고 '번호'가 들어간 줄을 열 제목으로 씁니다."""
    raw = pd.DataFrame(rows).fillna("").astype(str)
    if raw.empty:
        return raw
    header_row = 0
    for i in range(min(20, len(raw))):
        if any(any(k in str(v) for k in 제목_키워드) for v in raw.iloc[i].tolist()):
            header_row = i
            break
    names = []
    for idx, v in enumerate(raw.iloc[header_row].tolist()):
        name = str(v).strip() or f"열{idx + 1}"
        while name in names:
            name += "_"
        names.append(name)
    df = raw.iloc[header_row + 1:].copy()
    df.columns = names
    df = df.apply(lambda col: col.str.strip())
    df = df.replace(r"^(\d{4}-\d{2}-\d{2}) 00:00:00$", r"\1", regex=True)
    df = df[(df != "").any(axis=1)]
    df = df.loc[:, [c for c in df.columns if not (c.startswith("열") and (df[c] == "").all())]]
    return df.reset_index(drop=True)


@st.cache_data(show_spinner=False)
def read_ledger_file(file_bytes, file_name):
    """반환값: ({시트이름: 표}, 오류메시지)"""
    if file_bytes[:5] == b"SCDSA":
        return None, "DRM"
    try:
        if file_name.lower().endswith(".csv"):
            text = None
            for enc in ("utf-8-sig", "cp949"):
                try:
                    text = file_bytes.decode(enc)
                    break
                except UnicodeDecodeError:
                    pass
            if text is None:
                return None, "CSV 파일의 글자 형식을 알 수 없어요."
            rows = pd.read_csv(BytesIO(text.encode("utf-8")), header=None, dtype=str).fillna("").values.tolist()
            return {"CSV": rows_to_table(rows)}, None
        sheets = pd.read_excel(BytesIO(file_bytes), sheet_name=None, header=None, dtype=str)
        return {name: rows_to_table(df.fillna("").values.tolist()) for name, df in sheets.items()}, None
    except Exception as e:
        if "encrypted" in str(e).lower() or "not a zip file" in str(e).lower():
            return None, "DRM"
        return None, f"파일을 읽는 중 문제가 생겼어요: {e}"


def normalize_id(s):
    """대소문자, 띄어쓰기, 하이픈(-), 밑줄(_), 점(.)을 무시하고 비교하기 위한 변환"""
    return re.sub(r"[\s\-_.]", "", str(s).upper())


def find_match(job, candidate):
    """대장에서 번호를 찾습니다. 똑같은 게 없으면 띄어쓰기·하이픈 등을 무시하고 한 번 더 찾습니다."""
    candidate = candidate.strip()
    idx = job["columns"].index(job["id_col"])
    ids = [r[idx] for r in job["rows"]]
    if candidate in ids:
        return candidate
    target = normalize_id(candidate)
    hits = {v for v in ids if v and normalize_id(v) == target}
    return hits.pop() if target and len(hits) == 1 else None


def row_of(job, asset_id):
    idx = job["columns"].index(job["id_col"])
    for r in job["rows"]:
        if r[idx] == asset_id:
            return dict(zip(job["columns"], r))
    return {}


# ══════════════════════════════════════════════════════════════
# 사진 → AI로 자산번호 읽기
# ══════════════════════════════════════════════════════════════
def prepare_image(file_bytes):
    img = Image.open(BytesIO(file_bytes))
    img = ImageOps.exif_transpose(img)       # 휴대폰 사진이 옆으로 누워 보이는 문제 해결
    img = img.convert("RGB")
    img.thumbnail((1568, 1568))
    buf = BytesIO()
    img.save(buf, format="JPEG", quality=90)
    return base64.standard_b64encode(buf.getvalue()).decode("utf-8")


def read_asset_numbers(api_key, b64_image):
    """반환값: (후보 목록, 오류메시지)"""
    try:
        client = anthropic.Anthropic(api_key=api_key)
        response = client.beta.messages.create(
            model=MODEL,
            max_tokens=2000,
            output_config={"effort": "low"},     # 단순한 글자 읽기라 가볍게 처리 (비용 절약)
            betas=["server-side-fallback-2026-07-01"],
            fallbacks="default",                 # AI가 요청을 거절하면 다른 모델이 대신 처리
            messages=[{
                "role": "user",
                "content": [
                    {"type": "image", "source": {"type": "base64", "media_type": "image/jpeg", "data": b64_image}},
                    {"type": "text", "text": VISION_PROMPT},
                ],
            }],
        )
    except anthropic.AuthenticationError:
        return [], "API 키가 올바르지 않아요. 아래 'API 키 설정'에서 다시 입력해주세요."
    except anthropic.RateLimitError:
        return [], "요청이 너무 많아요. 잠시 후 다시 찍어주세요."
    except anthropic.APIConnectionError:
        return [], "인터넷 연결을 확인해주세요."
    except anthropic.BadRequestError as e:
        if "credit balance" in str(e).lower():
            return [], "API 크레딧(잔액)이 부족해요. console.anthropic.com에서 충전해주세요."
        return [], f"요청 오류: {e}"
    except anthropic.APIStatusError as e:
        return [], f"AI 서버 오류가 발생했어요 (코드: {e.status_code}). 잠시 후 다시 찍어주세요."
    except Exception as e:
        return [], f"사진을 처리하는 중 오류가 발생했어요: {e}"

    if response.stop_reason == "refusal":
        return [], None
    text = "".join(b.text for b in response.content if b.type == "text")
    candidates = []
    for line in text.splitlines():
        line = line.strip().strip("\"'`").strip()
        line = re.sub(r"^([-*•]|\d+[.)])\s+", "", line)
        if line and line.upper() != "NOT_FOUND":
            candidates.append(line)
    return candidates[:5], None


def check_candidates(job, candidates, source):
    """번호 후보들을 대장과 대조해서 기록하고, 화면에 보여줄 결과를 돌려줍니다."""
    if not candidates:
        return {"status": "not_found"}
    for cand in candidates:
        match = find_match(job, cand)
        if match:
            if match in job["checked"]:
                return {"status": "already", "id": match}
            job["checked"][match] = {"시각": now_text(), "방법": source}
            save_job(job)
            return {"status": "ok", "id": match}
    job["unknown"].append({"번호": candidates[0], "출처": source, "시각": now_text()})
    save_job(job)
    return {"status": "unknown", "id": candidates[0], "others": candidates[1:]}


# ══════════════════════════════════════════════════════════════
# 결과 엑셀 만들기
# ══════════════════════════════════════════════════════════════
def build_result_excel(job):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "실사결과"
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    ok_fill = PatternFill(start_color="DFF6DD", end_color="DFF6DD", fill_type="solid")
    miss_fill = PatternFill(start_color="FDE2E2", end_color="FDE2E2", fill_type="solid")

    ws.append(job["columns"] + ["실사결과", "확인시각", "확인방법"])
    for cell in ws[1]:
        cell.font, cell.fill = head_font, head_fill
    idx = job["columns"].index(job["id_col"])
    for r in job["rows"]:
        info = job["checked"].get(r[idx])
        ws.append(r + (["있음", info["시각"], info["방법"]] if info else ["미확인", "", ""]))
        ws.cell(row=ws.max_row, column=len(job["columns"]) + 1).fill = ok_fill if info else miss_fill
    for col_cells in ws.columns:
        width = max(len(str(c.value or "")) for c in col_cells[:200])
        ws.column_dimensions[col_cells[0].column_letter].width = min(max(10, width * 1.6), 50)

    if job["unknown"]:
        ws2 = wb.create_sheet("대장에없는번호")
        ws2.append(["읽은 번호", "사진/입력", "시각"])
        for cell in ws2[1]:
            cell.font, cell.fill = head_font, head_fill
        for item in job["unknown"]:
            ws2.append([item["번호"], item["출처"], item["시각"]])
        for letter in "ABC":
            ws2.column_dimensions[letter].width = 30

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ══════════════════════════════════════════════════════════════
# ① 자산대장 준비
# ══════════════════════════════════════════════════════════════
job = load_job(st.session_state.ai_job_id) if st.session_state.ai_job_id else None

if job is None:
    st.markdown("### ① 자산대장 준비")

    # 전에 하던 실사가 있으면 이어서 하기
    saved_jobs = list_jobs()
    if saved_jobs:
        labels = {j["id"]: f"{j['name']}  (확인 {len(j['checked'])}/{len(j['rows'])}, {j.get('updated', '')[:16]})"
                  for j in saved_jobs}
        pick = st.selectbox("전에 하던 실사 이어서 하기", list(labels.keys()), format_func=lambda k: labels[k])
        if st.button("▶ 이어서 하기", width="stretch", type="primary"):
            st.session_state.ai_job_id = pick
            st.session_state.ai_last = None
            st.rerun()
        st.caption("또는 아래에서 새 자산대장 파일을 올려주세요.")

    uploaded = st.file_uploader("자산대장 파일 올리기 (.xlsx 또는 .csv)", type=["xlsx", "xls", "csv"])
    if uploaded is not None:
        file_bytes = uploaded.getvalue()
        sheets, error = read_ledger_file(file_bytes, uploaded.name)
        if error == "DRM":
            st.error("🔒 회사 보안(DRM)으로 암호화된 파일이라 여기서는 읽을 수 없어요.")
            st.info("회사 PC에서 보안을 해제한 파일을 올려주세요. "
                    "외부에 나가도 괜찮은 열(자산번호, 자산명, 위치 등)만 남겨두는 것을 권장해요.")
        elif error:
            st.error(error)
        else:
            sheet_names = [n for n, df in sheets.items() if not df.empty]
            if not sheet_names:
                st.warning("파일에 자산 목록이 없어요.")
            else:
                sheet = st.selectbox("시트 선택", sheet_names) if len(sheet_names) > 1 else sheet_names[0]
                df = sheets[sheet]
                cols = list(df.columns)
                guess = next((c for k in 제목_키워드 for c in cols if k in c), cols[0])
                id_col = st.selectbox("어느 열이 '자산번호'인가요?", cols, index=cols.index(guess))
                st.caption(f"총 {len(df)}개 자산")
                st.dataframe(df.head(5), width="stretch", hide_index=True)

                if st.button("✅ 이 대장으로 실사 시작", width="stretch", type="primary"):
                    # 같은 파일·시트면 같은 기록 이름 → 전에 하던 기록이 있으면 이어짐
                    job_id = hashlib.md5(file_bytes + sheet.encode("utf-8")).hexdigest()[:16]
                    job = load_job(job_id) or {
                        "id": job_id, "name": uploaded.name, "checked": {}, "unknown": [],
                    }
                    job.update({"columns": cols, "rows": df.values.tolist(), "id_col": id_col})
                    save_job(job)
                    st.session_state.ai_job_id = job_id
                    st.session_state.ai_last = None
                    st.rerun()
    st.stop()


# ══════════════════════════════════════════════════════════════
# API 키 준비 (파일 → Streamlit Secrets 순서로 찾기)
# ══════════════════════════════════════════════════════════════
if "ai_api_key" not in st.session_state:
    key = get_key("anthropic_api_key")
    if not key:
        try:
            key = st.secrets.get("anthropic_api_key", "")
        except Exception:
            key = ""
    st.session_state.ai_api_key = key

id_idx = job["columns"].index(job["id_col"])
all_ids = [r[id_idx] for r in job["rows"] if r[id_idx]]
total = len(all_ids)
done = sum(1 for v in all_ids if v in job["checked"])

top_l, top_r = st.columns([3, 1])
top_l.markdown(f"**📄 {job['name']}**  \n확인 **{done} / {total}** ({done / total * 100 if total else 0:.0f}%)")
if top_r.button("대장 바꾸기"):
    st.session_state.ai_job_id = None
    st.session_state.ai_last = None
    st.rerun()
st.progress(done / total if total else 0)

# ══════════════════════════════════════════════════════════════
# ② 촬영 → 바로 인식
# ══════════════════════════════════════════════════════════════
st.markdown("### ② 자산번호 촬영")

if not st.session_state.ai_api_key:
    st.warning("사진 인식을 쓰려면 맨 아래 **🔑 API 키 설정**에서 Claude API 키를 입력해주세요. "
               "(키 없이도 '번호 직접 입력'은 쓸 수 있어요)")

mode = st.radio("촬영 방식", ["📱 휴대폰 카메라 (고화질·추천)", "🎥 화면 안에서 찍기"],
                horizontal=True, label_visibility="collapsed")
photo_key = f"ai_photo_{st.session_state.ai_photo_key}"
if mode.startswith("📱"):
    st.caption("아래 **Browse files / 파일 선택**을 누르고 **'사진 찍기(카메라)'**를 고르세요. 찍자마자 자동으로 확인돼요.")
    photo = st.file_uploader("📷 자산번호 라벨 촬영", type=["jpg", "jpeg", "png", "heic", "heif", "webp"],
                             key=photo_key, label_visibility="collapsed")
else:
    photo = st.camera_input("📷 자산번호 라벨을 화면 가운데에 맞추고 찍어주세요", key=photo_key)

if photo is not None and st.session_state.ai_api_key:
    photo_bytes = photo.getvalue()
    with st.spinner("AI가 자산번호를 읽고 있어요..."):
        try:
            b64 = prepare_image(photo_bytes)
            candidates, error = read_asset_numbers(st.session_state.ai_api_key, b64)
        except Exception as e:
            candidates, error = [], f"사진을 열 수 없어요. 다른 사진으로 다시 찍어주세요. ({e})"
    if error:
        result = {"status": "error", "message": error}
    else:
        result = check_candidates(job, candidates, f"사진 {now_text()[11:16]}")
    # 작은 미리보기 사진도 같이 보관
    try:
        thumb = ImageOps.exif_transpose(Image.open(BytesIO(photo_bytes)))
        thumb.thumbnail((400, 400))
        buf = BytesIO()
        thumb.convert("RGB").save(buf, format="JPEG", quality=80)
        result["thumb"] = buf.getvalue()
    except Exception:
        pass
    st.session_state.ai_last = result
    st.session_state.ai_photo_key += 1      # 촬영 칸을 비워서 바로 다음 자산을 찍을 수 있게 함
    st.rerun()
elif photo is not None:
    st.error("API 키가 없어서 사진을 읽을 수 없어요. 맨 아래 🔑 API 키 설정을 먼저 해주세요.")


def show_row_details(asset_id):
    info = row_of(job, asset_id)
    details = [f"{k}: {v}" for k, v in info.items() if k != job["id_col"] and v]
    if details:
        st.caption(" · ".join(details))


# 마지막 인식 결과 크게 보여주기
last = st.session_state.ai_last
if last:
    status = last["status"]
    if status == "ok":
        st.markdown(f"<div style='background:#DFF6DD;border-radius:14px;padding:18px;text-align:center;"
                    f"font-size:1.5em;font-weight:700;color:#0F7B0F'>✅ 있음<br>{last['id']}</div>",
                    unsafe_allow_html=True)
        show_row_details(last["id"])
    elif status == "already":
        st.info(f"🔁 **{last['id']}** 는 이미 확인된 자산이에요.")
        show_row_details(last["id"])
    elif status == "unknown":
        st.markdown(f"<div style='background:#FDE2E2;border-radius:14px;padding:18px;text-align:center;"
                    f"font-size:1.3em;font-weight:700;color:#C42B1C'>❌ 대장에 없는 번호<br>{last['id']}</div>",
                    unsafe_allow_html=True)
        if last.get("others"):
            st.caption("AI가 읽은 다른 후보: " + ", ".join(last["others"]))
        st.caption("AI가 잘못 읽었다면 아래 '번호 직접 입력'에 올바른 번호를 넣어주세요. (이 번호는 '대장에 없는 번호'로 기록돼요)")
    elif status == "not_found":
        st.warning("⚠ 사진에서 자산번호를 읽지 못했어요. 라벨에 더 가까이, 밝은 곳에서 다시 찍어주세요.")
    elif status == "error":
        st.error(last["message"])
    if last.get("thumb"):
        with st.expander("방금 찍은 사진 보기"):
            st.image(last["thumb"])

# 번호 직접 입력
with st.expander("✏️ 번호 직접 입력 (사진 없이 확인)", expanded=(last or {}).get("status") in ("unknown", "not_found")):
    with st.form("manual_form", clear_on_submit=True):
        manual = st.text_input("자산번호", placeholder="예: IT-2023-0015")
        if st.form_submit_button("확인", width="stretch") and manual.strip():
            st.session_state.ai_last = check_candidates(job, [manual.strip()], "직접 입력")
            st.rerun()

st.divider()

# ══════════════════════════════════════════════════════════════
# ③ 진행 현황 / 결과 저장
# ══════════════════════════════════════════════════════════════
st.markdown("### ③ 진행 현황")
c1, c2, c3 = st.columns(3)
c1.metric("확인", f"{done}건")
c2.metric("미확인", f"{total - done}건")
c3.metric("대장에 없음", f"{len(job['unknown'])}건")

df_all = pd.DataFrame(job["rows"], columns=job["columns"])
checked_mask = df_all[job["id_col"]].isin(job["checked"].keys())

with st.expander(f"아직 확인 안 된 자산 ({(~checked_mask).sum()}건)"):
    st.dataframe(df_all[~checked_mask], width="stretch", hide_index=True)

with st.expander(f"확인된 자산 ({checked_mask.sum()}건) — 잘못 표시된 건 여기서 취소"):
    if job["checked"]:
        view = pd.DataFrame([{"자산번호": k, "확인시각": v["시각"], "방법": v["방법"]}
                             for k, v in job["checked"].items()]).sort_values("확인시각", ascending=False)
        st.dataframe(view, width="stretch", hide_index=True)
        cancel = st.selectbox("취소할 자산번호", list(view["자산번호"]))
        if st.button("↩ 있음 표시 취소"):
            job["checked"].pop(cancel, None)
            save_job(job)
            st.session_state.ai_last = None
            st.rerun()
    else:
        st.caption("아직 없어요.")

if job["unknown"]:
    with st.expander(f"대장에 없는 번호 ({len(job['unknown'])}건)"):
        st.dataframe(pd.DataFrame(job["unknown"]), width="stretch", hide_index=True)

st.download_button(
    "📥 결과 엑셀 저장", data=build_result_excel(job),
    file_name=f"{os.path.splitext(job['name'])[0]}_실사결과_{datetime.now(KST).strftime('%Y%m%d')}.xlsx",
    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    width="stretch", type="primary",
)

st.divider()

# ── 설정 ──────────────────────────────────────────────────────
with st.expander("🔑 API 키 설정"):
    st.markdown("사진에서 자산번호를 읽으려면 Claude API 키가 필요해요. "
                "[console.anthropic.com](https://console.anthropic.com) → Settings → API Keys에서 발급받을 수 있어요.")
    new_key = st.text_input("API 키", value=st.session_state.ai_api_key, type="password")
    if st.button("💾 API 키 저장"):
        st.session_state.ai_api_key = new_key.strip()
        set_key("anthropic_api_key", new_key.strip())
        st.success("저장했어요.")

with st.expander("⚙️ 자산번호 열 바꾸기 / 기록 초기화"):
    new_col = st.selectbox("자산번호 열", job["columns"], index=job["columns"].index(job["id_col"]))
    if new_col != job["id_col"] and st.button("열 바꾸기"):
        job["id_col"] = new_col
        save_job(job)
        st.rerun()
    confirm = st.checkbox("이 대장의 실사 기록(있음 표시)을 모두 지울게요")
    if st.button("🗑 기록 초기화", disabled=not confirm):
        job["checked"], job["unknown"] = {}, []
        save_job(job)
        st.session_state.ai_last = None
        st.rerun()
